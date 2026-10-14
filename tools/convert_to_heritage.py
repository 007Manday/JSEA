"""Convert MM-4034 JSEA documents into the Heritage Minerals JSEA/SWMS form.

Usage: python tools/convert_to_heritage.py TEMPLATE.docx SRC_DIR OUT_DIR

TEMPLATE is a filled Heritage JSEA/SWMS form (HM-MMM-WHS-FRM-031-R03). Its
header block, risk matrix, SOP list and sign-on pages are kept; the work
activity details, ticks and job-step rows are replaced from each source JSEA.

Source P x C is carried as the Initial Risk. Residual Risk is estimated as
likelihood reduced by two levels (minimum 1) with consequence unchanged, and
rated against the Heritage 5x5 matrix.
"""
import copy
import glob
import os
import re
import sys
from datetime import datetime

import docx
from docx.oxml.ns import qn

# Heritage 5x5 matrix, [likelihood][consequence]
MATRIX = {
    5: ['M', 'H', 'H', 'E', 'E'],
    4: ['M', 'M', 'H', 'H', 'E'],
    3: ['L', 'M', 'M', 'H', 'H'],
    2: ['L', 'L', 'M', 'M', 'H'],
    1: ['L', 'L', 'L', 'M', 'M'],
}


def rate(l, c):
    return f'{MATRIX[l][c - 1]}{l * c}'


def uniq_cells(row):
    seen = []
    for c in row.cells:
        if all(c._tc is not s._tc for s in seen):
            seen.append(c)
    return seen


def _run_props(p):
    r = p.find(qn('w:r'))
    if r is not None and r.find(qn('w:rPr')) is not None:
        return copy.deepcopy(r.find(qn('w:rPr')))
    ppr = p.find(qn('w:pPr'))
    if ppr is not None and ppr.find(qn('w:rPr')) is not None:
        return copy.deepcopy(ppr.find(qn('w:rPr')))
    return None


def _fill_para(p, text):
    rpr = _run_props(p)
    for child in list(p):
        if child.tag != qn('w:pPr'):
            p.remove(child)
    r = p.makeelement(qn('w:r'), {})
    if rpr is not None:
        r.append(rpr)
    t = r.makeelement(qn('w:t'), {})
    t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.append(t)
    p.append(r)


def set_lines(cell, lines):
    """Fill a cell with one paragraph per line, cloned from its first paragraph."""
    tc = cell._tc
    paras = tc.findall(qn('w:p'))
    proto = copy.deepcopy(paras[0])
    for p in paras:
        tc.remove(p)
    for line in lines or ['']:
        p = copy.deepcopy(proto)
        _fill_para(p, line)
        tc.append(p)


def set_text(cell, text, size=None):
    set_lines(cell, [text])
    if size:
        for rpr in cell._tc.iter(qn('w:rPr')):
            for tag in ('w:sz', 'w:szCs'):
                el = rpr.find(qn(tag))
                if el is not None:
                    el.set(qn('w:val'), str(size))


def set_check(cell, on):
    for cb in cell._tc.iter(qn('w:checkBox')):
        for ch in cb.findall(qn('w:checked')):
            cb.remove(ch)
        cb.find(qn('w:default')).set(qn('w:val'), '1' if on else '0')


def yes_no(cells, yes):
    set_check(cells[0], yes)
    set_check(cells[1], not yes)


def section_lines(doc, start, stop):
    out, on = [], False
    for p in doc.paragraphs:
        t = p.text.strip()
        if re.match(rf'^{start}\s', t):
            on = True
            continue
        if on and re.match(rf'^{stop}\s', t):
            break
        if on and t:
            out.append(t)
    return out


def sentences(text):
    out = []
    for part in re.split(r'(?<=\.)\s+(?=[A-Z])|\n', text):
        part = part.strip().rstrip('.')
        if part and part not in out:
            out.append(part)
    return out


def read_source(path):
    d = docx.Document(path)
    meta = {r.cells[0].text.strip(): r.cells[1].text.strip() for r in d.tables[0].rows}
    title = d.paragraphs[1].text.strip()
    risk = next(t for t in d.tables if t.rows[0].cells[0].text == 'Job step')
    steps = []
    for r in risk.rows[1:]:
        c = r.cells
        steps.append({
            'step': re.sub(r'\s+', ' ', c[0].text.strip()),
            'hazards': [p.text.strip() for p in c[1].paragraphs if p.text.strip()],
            'controls': sentences(c[2].text),
            'p': int(c[3].text), 'c': int(c[4].text),
            'who': c[6].text.strip() or 'Operator',
        })
    ppe = section_lines(d, '5', '6')
    training = section_lines(d, '8', '9')
    full = '\n'.join(p.text for p in d.paragraphs) + '\n' + '\n'.join(
        c.text for t in d.tables for r in t.rows for c in r.cells)
    return {'meta': meta, 'title': title, 'steps': steps, 'ppe': ppe,
            'training': training[1:], 'text': full}


def has(text, *words):
    low = text.lower()
    return any(w in low for w in words)


def hierarchy(controls):
    t = ' '.join(controls)
    return [
        has(t, 'eliminat'),
        has(t, 'substitut'),
        has(t, 'isolat', 'lock-out', 'lockout', 'loto'),
        has(t, 'guard', 'interlock', 'fixed detection', 'handrail', 'fixed platform',
            'ventilat', 'extraction', 'fume cupboard', 'bunded', 'enclosed'),
        True,
        has(t, 'glove', 'goggle', 'face shield', 'monitor', 'harness', 'apron',
            'respirator', 'gas mask', 'boots', 'ppe', 'hearing'),
    ]


CHEMICALS = [
    (('cyanide', 'nacn'), 'Sodium cyanide (NaCN) / cyanide-bearing slurry and solution'),
    (('caustic', 'naoh', 'sodium hydroxide'), 'Caustic soda (NaOH)'),
    (('sulphuric', 'h2so4'), 'Sulphuric acid (H2SO4)'),
    (('hydrochloric', ' hcl'), 'Hydrochloric acid (HCl)'),
    (('silver nitrate', 'agno3'), 'Silver nitrate (AgNO3) titrant'),
    (('lime', 'cao'), 'Lime / lime slurry'),
    (('sodium chloride', 'nacl', 'brine'), 'Sodium chloride (NaCl) brine'),
    (('flocculant',), 'Flocculant'),
    (('copper sulphate', 'cuso4'), 'Copper sulphate'),
    (('sodium metabisulphite', 'smbs'), 'Sodium metabisulphite (SMBS)'),
]


def chemicals(text):
    found = [name for keys, name in CHEMICALS if has(text, *keys)]
    if len(found) > 3:
        found = found[:2] + ['; '.join(found[2:])]
    return found


def convert(template, src_path, out_path):
    s = read_source(src_path)
    m, text = s['meta'], s['text']
    ppe_text = ' '.join(s['ppe'])
    doc = docx.Document(template)
    t0 = doc.tables[0]
    R = lambda i: uniq_cells(t0.rows[i])

    # Header block
    date = datetime.strptime(m['Date'], '%d %B %Y').strftime('%d-%m-%Y')
    num = m['Document number']
    set_text(R(0)[5], date)
    set_text(R(1)[3], '')
    set_text(R(1)[5], '12 Months')
    set_text(R(2)[1], f"{s['title']} ({m['Area']})")
    set_text(R(2)[3], f'HM-MMM-WHS-FRM-031-R03-{num}')

    # PPE ticks: boots, gloves, hard hat & glasses, hearing, hi-vis, dust mask,
    # harness, isolation, hand washing, double eye protection
    ppe = [True, True, True,
           has(ppe_text, 'hearing', 'ear plug', 'earmuff'),
           True,
           has(ppe_text, 'mask', 'respirator', 'dust'),
           has(ppe_text, 'harness'),
           has(text, 'isolation', 'lock-out', 'lockout', 'loto'),
           has(text, 'cyanide', 'chemical', 'acid', 'caustic', 'lead', 'mercury'),
           has(ppe_text, 'goggles and face shield', 'face shield')]
    for cell, on in zip(R(5), ppe):
        set_check(cell, on)

    # Hazardous chemicals (rows 12-14)
    chems = chemicals(text)
    for i, ri in enumerate((12, 13, 14)):
        c = R(ri)
        present = i < len(chems)
        set_text(c[0], chems[i] if present else '')
        for a, b in ((1, 2), (3, 4), (6, 7)):
            if present:
                yes_no((c[a], c[b]), True)
            else:
                set_check(c[a], False)
                set_check(c[b], False)
        set_text(c[5], '')

    # High risk work (rows 16-19: label/checkbox pairs)
    hrw = {
        'Excavation / Trenching': False,
        'Falling from Height': has(text, 'fall from', 'at height'),
        'Rigging / Lifting': has(text, 'crane', 'rigging'),
        'Powered Mobile Plant': has(text, 'forklift', 'mobile plant'),
        'Scaffolding': has(text, 'scaffold'),
        'Pressurised Equipment': has(text, 'pressurised', 'under pressure', 'pressure line'),
        'Confined Space': has(text, 'confined space'),
        'Chemical/Fuel/Refridge line': has(text, 'cyanide', 'acid', 'caustic'),
        'Extremes of Temperature': has(text, 'heat stress', 'hot solution', 'burn', 'scald'),
        'Traffic Corridors': has(text, 'traffic', 'light vehicle'),
        'Electrical': has(text, 'electrical', 'electrocution'),
        'Risk of Drowning': has(text, 'drown'),
    }
    for ri in (16, 17, 18, 19):
        c = R(ri)
        for k in range(0, len(c), 2):
            set_check(c[k + 1], hrw.get(c[k].text.strip(), False))

    # Permits (rows 22-23)
    c = R(22)
    yes_no((c[1], c[2]), False)
    yes_no((c[5], c[6]), has(ppe_text, 'harness'))
    c = R(23)
    yes_no((c[1], c[2]), has(text, 'confined space entry', 'confined space permit'))
    yes_no((c[5], c[6]), has(text, 'hot work'))

    # Training (rows 26-30)
    c = R(26); yes_no((c[3], c[4]), True)
    c = R(27); yes_no((c[3], c[4]), True)
    other = [t.rstrip('.') for t in s['training'] if not has(t, 'site induction')][:3]
    for i, ri in enumerate((28, 29, 30)):
        c = R(ri)
        if i < len(other):
            set_text(c[3], other[i])
            yes_no((c[4], c[5]), True)
        else:
            set_text(c[3], '')
            set_check(c[4], False)
            set_check(c[5], False)

    # Heritage SOPs (rows 33-64: name/checkbox pairs)
    sops = {
        'Risk Management': True,
        'Incident Reporting': True,
        'Emergency Response': True,
        'Induction & Training': True,
        'Permit': has(ppe_text, 'harness') or has(text, 'permit to work', 'confined space entry'),
        'Working in Isolation': has(text, 'alone'),
        'Working at Heights': hrw['Falling from Height'],
        'Confined Space': hrw['Confined Space'],
        'Isolation': ppe[7],
        'Plant & Equipment': has(text, 'pump', 'agitator', 'screen', 'cutter'),
        'Electrical': hrw['Electrical'],
        'Fatigue': has(text, 'fatigue', 'night shift'),
        'Hazardous Substances': bool(chems),
        'Waste Management': True,
        'Water Management': has(text, 'spill', 'bund', 'drain', 'return water'),
        'Dust & Air Management': has(text, 'dust', 'hcn', 'fume'),
        'Monitoring & Measurement': True,
        'Record Management': True,
        'non-conformance': has(text, 'out of range', 'out-of-range', 'non-conform'),
    }
    for ri in range(33, 65):
        c = R(ri)
        for k in range(0, len(c) - 1, 2):
            name = c[k].text.strip()
            if 'SOP' not in name:
                continue
            key = re.sub(r'^.*?-R02-', '', name)
            set_check(c[k + 1], sops.get(key, False))

    # Job steps
    t2 = doc.tables[2]
    proto = copy.deepcopy(t2.rows[4]._tr)
    for row in list(t2.rows)[4:]:
        t2._tbl.remove(row._tr)
    for st in s['steps']:
        tr = copy.deepcopy(proto)
        t2._tbl.append(tr)
        cells = uniq_cells(t2.rows[-1])
        assert len(cells) == 16, len(cells)
        rl, rc = max(1, st['p'] - 2), st['c']
        set_text(cells[0], st['step'])
        set_lines(cells[1], st['hazards'])
        set_text(cells[2], str(st['p']))
        set_text(cells[3], str(st['c']))
        set_text(cells[4], rate(st['p'], st['c']), size=16)
        set_lines(cells[5], st['controls'])
        set_text(cells[6], str(rl))
        set_text(cells[7], str(rc))
        set_text(cells[8], rate(rl, rc), size=16)
        for cell, on in zip(cells[9:15], hierarchy(st['controls'])):
            set_text(cell, 'x' if on else '')
        set_text(cells[15], st['who'])

    doc.core_properties.title = f"{num} {s['title']}"
    doc.save(out_path)


def main():
    template, src_dir, out_dir = sys.argv[1:4]
    os.makedirs(out_dir, exist_ok=True)
    for src in sorted(glob.glob(os.path.join(src_dir, '*.docx'))):
        name = os.path.basename(src)
        name = re.sub(r'^MM-4034-033-OR-03-', '', name)
        name = re.sub(r'-Rev0\.docx$', '.docx', name)
        out = os.path.join(out_dir, name)
        convert(template, src, out)
        print(out)


if __name__ == '__main__':
    main()
